import os
import sys
import re
import json
import glob
import time
import unicodedata
from groq import Groq
from pypdf import PdfReader
from docx import Document

# Asegurar codificación UTF-8 en consola de Windows
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# 1. Configuración de Groq API y lectura de .env
def load_env_file(filepath=".env"):
    """Carga variables del archivo .env si existen y no están en el entorno."""
    if os.path.exists(filepath):
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith("#") and "=" in line:
                        k, v = line.split("=", 1)
                        k = k.strip()
                        v = v.strip().strip('"').strip("'")
                        if k not in os.environ:
                            os.environ[k] = v
        except Exception as e:
            print(f"[!] Aviso al leer .env: {e}")

load_env_file()

api_key = os.environ.get("GROQ_API_KEY")
if not api_key:
    print("[-] ERROR: GROQ_API_KEY no encontrada en variables de entorno ni en .env")
    sys.exit(1)

client = Groq(api_key=api_key)

RAW_FOLDER = "entradas_crudos"
OUTPUT_FILE = os.path.join("src", "data", "news.json")
MODEL_NAME = "llama-3.3-70b-versatile"
FALLBACK_MODEL = "llama-3.1-8b-instant"

def slugify(text: str) -> str:
    """Convierte texto en un slug kebab-case limpio."""
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^\w\s-]", "", text.lower()).strip()
    return re.sub(r"[-\s]+", "-", text)

def parse_filename_metadata(filename: str):
    """
    Extrae autor y fecha desde el nombre de archivo con formato:
    'Nombre del autor DD-MM-AAAA.ext' o 'Nombre del autor DD-MM-AA.ext'
    Retorna (autor_sugerido, fecha_iso_aaaa_mm_dd)
    """
    name_without_ext = os.path.splitext(filename)[0].strip()
    
    # Busca el patrón DD-MM-AA o DD-MM-AAAA al final del nombre
    date_match = re.search(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})\s*$", name_without_ext)
    
    suggested_author = None
    suggested_date_iso = None
    
    if date_match:
        day, month, year = date_match.groups()
        if len(year) == 2:
            year = f"20{year}"
        suggested_date_iso = f"{year}-{int(month):02d}-{int(day):02d}"
        
        # El resto del nombre previo a la fecha es el autor
        author_part = name_without_ext[:date_match.start()].strip(" _-")
        if author_part:
            suggested_author = re.sub(r"\s+", " ", author_part.replace("_", " ")).strip()
    else:
        suggested_author = re.sub(r"\s+", " ", name_without_ext.replace("_", " ")).strip()
        
    return suggested_author, suggested_date_iso

def extract_text_from_pdf(pdf_path: str) -> str:
    """Extrae texto completo de un archivo PDF."""
    text = ""
    try:
        reader = PdfReader(pdf_path)
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text += content + "\n"
    except Exception as e:
        print(f"[-] Error al leer PDF {pdf_path}: {e}")
    return text.strip()

def extract_text_from_docx(docx_path: str) -> str:
    """Extrae texto completo de un archivo Word (.docx), incluyendo párrafos y tablas."""
    text_parts = []
    try:
        doc = Document(docx_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_parts.append(row_text)
    except Exception as e:
        print(f"[-] Error al leer DOCX {docx_path}: {e}")
    return "\n\n".join(text_parts).strip()

def standardize_with_groq(raw_text: str, filename: str, fallback_author: str, fallback_date: str) -> dict:
    """Envía el texto a Groq indicándole los metadatos de respaldo del archivo con reintentos."""
    system_prompt = f"""
    Eres un editor académico y analista senior del Observatorio de Economía Internacional y Relaciones Internacionales de la carrera de Licenciatura en Relaciones Internacionales (Materia: Economía IV).
    Tu tarea es procesar el texto crudo del trabajo práctico semanal de un estudiante sobre una noticia económica y devolver un objeto JSON estandarizado para la plataforma web.

    METADATOS DE RESPALDO OBTENIDOS DEL NOMBRE DE ARCHIVO:
    - Autor por defecto: "{fallback_author or 'Estudiante de Relaciones Internacionales'}"
    - Fecha por defecto (ISO YYYY-MM-DD): "{fallback_date or '2026-09-02'}"

    REGLAS OBLIGATORIAS:
    1. Si el autor NO está explícito en el texto, UTILIZA OBLIGATORIAMENTE el "Autor por defecto".
    2. Si la fecha NO se indica claramente en el texto, UTILIZA OBLIGATORIAMENTE la "Fecha por defecto".
    3. Siempre entrega la fecha en formato ISO 'YYYY-MM-DD'.
    4. "region" debe ser una de: "América Latina", "Global", "Europa", "Asia-Pacífico", "Norteamérica".
    5. "topic" debe ser un área temática económica clara (ej: "Política Monetaria", "Comercio Exterior", "Integración Regional", "Energía & Recursos", "Monedas & Deuda", "Empleo & Producción", "Macroeconomía").
    6. "authors" DEBE ser una lista de strings con nombres de los estudiantes autores (ej: ["{fallback_author or 'Estudiante'}"]). NUNCA una lista de objetos.
    7. "summary_bullets" debe ser una lista de 2 a 4 puntos clave sintéticos y directos.
    8. "full_text" DEBE ser una lista de 3 a 5 párrafos completos y bien redactados con el desarrollo analítico de la noticia y sus implicancias económicas.
    9. "glossary" debe contener de 2 a 5 términos económicos relevantes con su definición pedagógica clara.
    10. "ir_perspective" debe contener:
        - "geopolitical_context": Análisis del impacto geopolítico, dinámicas de poder e interdependencia internacional.
        - "strategic_recommendations": Lista de 2 o 3 recomendaciones de política exterior o estrategia comercial.
    11. "sources" debe ser una lista con el medio o fuentes citadas (ej: ["ABC Color", "Banco Central del Paraguay (BCP)"]).

    Responde EXCLUSIVAMENTE con un JSON válido con la siguiente estructura exacta:
    {{
      "id": "slug-identificador-en-kebab-case",
      "title": "Título analítico y periodístico de la noticia",
      "region": "Región",
      "topic": "Tema",
      "date": "YYYY-MM-DD",
      "authors": ["Nombre Completo"],
      "author_role": "Estudiante Investigador • Relaciones Internacionales",
      "author_image": "/authors/sofia-valenzuela.jpg",
      "summary_bullets": [
        "Punto clave 1...",
        "Punto clave 2..."
      ],
      "full_text": [
        "Párrafo 1 de análisis...",
        "Párrafo 2 con datos y contexto...",
        "Párrafo 3 con conclusiones..."
      ],
      "glossary": [
        {{
          "term": "Concepto",
          "definition": "Definición pedagógica"
        }}
      ],
      "ir_perspective": {{
        "geopolitical_context": "Análisis geopolítico y de impacto exterior...",
        "strategic_recommendations": [
          "Recomendación estratégica 1...",
          "Recomendación estratégica 2..."
        ]
      }},
      "sources": [
        "Medio o fuente original"
      ]
    }}
    """

    models_to_try = [MODEL_NAME, FALLBACK_MODEL]
    for model_candidate in models_to_try:
        max_retries = 4
        for attempt in range(max_retries):
            try:
                response = client.chat.completions.create(
                    model=model_candidate,
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": f"Archivo: {filename}\n\nContenido del trabajo del alumno:\n{raw_text[:6000]}"}
                    ],
                    response_format={"type": "json_object"},
                    temperature=0.1
                )
                data = json.loads(response.choices[0].message.content)
                return data
            except Exception as e:
                err_msg = str(e)
                if "rate_limit" in err_msg.lower() or "429" in err_msg:
                    if model_candidate != models_to_try[-1] and attempt == 0:
                        print(f"    [⚡ Rate limit en {model_candidate}] Alternando al modelo '{FALLBACK_MODEL}'...")
                        break
                    wait_time = 15 + attempt * 10
                    print(f"    [⏳ Rate limit en {model_candidate}] Reintentando en {wait_time}s (intento {attempt + 1}/{max_retries})...")
                    time.sleep(wait_time)
                elif attempt < max_retries - 1:
                    print(f"    [!] Error en Groq: {err_msg[:80]}. Reintentando en 3s...")
                    time.sleep(3)
                else:
                    if model_candidate != models_to_try[-1]:
                        print(f"    [↳ Alternando al modelo de respaldo '{FALLBACK_MODEL}']...")
                        break
                    print(f"[-] Error definitivo en Groq al procesar '{filename}': {e}")
                    return None

def main():
    if not os.path.exists(RAW_FOLDER):
        os.makedirs(RAW_FOLDER)
        print(f"[!] Carpeta '{RAW_FOLDER}' creada. Coloca allí los archivos.")
        return

    files = glob.glob(os.path.join(RAW_FOLDER, "*.pdf")) + glob.glob(os.path.join(RAW_FOLDER, "*.docx"))
    if not files:
        print(f"[!] No se encontraron archivos .pdf ni .docx en '{RAW_FOLDER}'.")
        return

    # Cargar base de datos existente de noticias
    existing_news = []
    if os.path.exists(OUTPUT_FILE):
        try:
            with open(OUTPUT_FILE, "r", encoding="utf-8") as f:
                existing_news = json.load(f)
                if not isinstance(existing_news, list):
                    existing_news = []
        except Exception:
            existing_news = []

    # Mapear IDs y archivos ya procesados para evitar duplicación y permitir reanudar
    existing_ids = {item.get("id") for item in existing_news if isinstance(item, dict) and item.get("id")}
    processed_files = {item.get("source_file") for item in existing_news if isinstance(item, dict) and item.get("source_file")}
    
    processed_count = 0
    total_files = len(files)

    print(f"[*] Iniciando procesamiento de {total_files} archivos...")
    print(f"[*] Artículos existentes en {OUTPUT_FILE}: {len(existing_news)}")

    for idx, file_path in enumerate(files, 1):
        filename = os.path.basename(file_path)
        
        # Saltar archivos ya procesados previamente
        if filename in processed_files:
            print(f"[{idx}/{total_files}] Omitido (ya procesado): {filename}")
            continue

        fallback_author, fallback_date = parse_filename_metadata(filename)
        
        print(f"\n[{idx}/{total_files}] [+] Procesando: {filename}")
        print(f"    ↳ Metadatos en nombre -> Autor: '{fallback_author}' | Fecha: '{fallback_date}'")

        if file_path.lower().endswith(".pdf"):
            raw_text = extract_text_from_pdf(file_path)
        elif file_path.lower().endswith(".docx"):
            raw_text = extract_text_from_docx(file_path)
        else:
            continue

        if not raw_text.strip():
            print(f"[-] Archivo vacío o ilegible: {filename}")
            continue

        item_data = standardize_with_groq(raw_text, filename, fallback_author, fallback_date)

        if item_data and "title" in item_data:
            # 1. Normalización de fecha
            if not item_data.get("date") or item_data.get("date") in ["YYYY-MM-DD", ""]:
                if fallback_date:
                    item_data["date"] = fallback_date

            # 2. Normalización de autores a list[str]
            raw_authors = item_data.get("authors")
            clean_authors = []
            if isinstance(raw_authors, list):
                for a in raw_authors:
                    if isinstance(a, str) and a.strip() and a.strip() not in ["Nombre Completo", "Nombre del alumno"]:
                        clean_authors.append(a.strip())
                    elif isinstance(a, dict) and a.get("name") and a.get("name") not in ["Nombre Completo", "Nombre del alumno"]:
                        clean_authors.append(a.get("name").strip())
            
            if not clean_authors:
                clean_authors = [fallback_author or "Estudiante de Relaciones Internacionales"]
            item_data["authors"] = clean_authors

            # 3. Metadatos editoriales
            if not item_data.get("author_role"):
                item_data["author_role"] = "Estudiante Investigador • Relaciones Internacionales"
            if not item_data.get("author_image"):
                item_data["author_image"] = "/authors/sofia-valenzuela.jpg"

            # 4. Asegurar campos obligatorios de la UI
            if not item_data.get("topic"):
                item_data["topic"] = item_data.get("category") or "Economía Internacional"
            if not item_data.get("region"):
                item_data["region"] = "América Latina"
            if not item_data.get("summary_bullets") or not isinstance(item_data.get("summary_bullets"), list):
                item_data["summary_bullets"] = [item_data.get("summary") or item_data.get("title")]
            if not item_data.get("full_text") or not isinstance(item_data.get("full_text"), list):
                item_data["full_text"] = [raw_text[:1200]]
            if not item_data.get("sources") or not isinstance(item_data.get("sources"), list):
                item_data["sources"] = [item_data.get("originalSource") or "Medio de Prensa"]

            # 5. Generación de ID único y limpio
            base_slug = item_data.get("id") or slugify(item_data.get("title") or "noticia")
            base_slug = slugify(base_slug)
            date_clean = (item_data.get("date") or "fecha").replace("-", "")
            unique_id = f"{base_slug}-{date_clean}"
            
            counter = 1
            temp_id = unique_id
            while temp_id in existing_ids:
                temp_id = f"{unique_id}-{counter}"
                counter += 1

            item_data["id"] = temp_id
            existing_ids.add(temp_id)

            # 6. Registrar archivo de origen
            item_data["source_file"] = filename
            processed_files.add(filename)

            existing_news.append(item_data)
            processed_count += 1
            
            print(f"[✓] Guardado: '{item_data.get('title')}' ({item_data.get('date')}) por {item_data['authors'][0]}")

            # Guardado progresivo tras cada archivo para no perder avances
            # Ordenar por fecha descendente
            existing_news.sort(key=lambda x: str(x.get("date", "")), reverse=True)
            os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
            with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
                json.dump(existing_news, f, ensure_ascii=False, indent=2)

        # Pausa suave de cortesía para respetar límites de tokens por minuto (TPM)
        time.sleep(1.2)

    print(f"\n[🚀] Ingesta completada con éxito.")
    print(f"Nuevos artículos procesados en esta sesión: {processed_count}")
    print(f".Total acumulado en catálogo ({OUTPUT_FILE}): {len(existing_news)} artículos.")

if __name__ == "__main__":
    main()